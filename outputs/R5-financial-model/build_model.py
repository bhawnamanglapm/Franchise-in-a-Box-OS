"""Build the FreshBrew financial model (Task 5) as an .xlsx with live formulas.

Every number the model uses is an input on the Assumptions tab. All other cells are formulas, so the
workbook recalculates when an input changes, in Excel or after import into Google Sheets.

Run:  python build_model.py            -> writes FreshBrew-Financial-Model.xlsx next to this file
      python build_model.py --expect   -> prints the values the formulas should produce (for checks)
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("FreshBrew-Financial-Model.xlsx")
ACCESSED = "accessed 17 Oct 2026"

# Sources. Third-party listings and trade articles, not franchisor disclosures; labelled as such.
S1 = ("S1: FranchiseBazar, 'Top Cafe Franchises Between 30 To 50 Lakhs in India 2026' (24 Jun 2026): "
      "Barista ₹30–50L, fee ₹5–8L, 8% royalty, 650–850 sq ft, break-even 24–36 months; "
      "BrewBakes fee ₹10L, 5% royalty; Brewtopia ₹30–50L, 400–1,000 sq ft")
S2 = ("S2: FranchiseBazar, 'Best Chai Franchise Business in India' (Aug 2026): "
      "Chai Sutta Bar ₹20–30L, 4% royalty; MBA Chai Wala ₹14–23L, ~4–6% royalty")
S3 = ("S3: Restaurant India, 'Can Restaurants Stay Profitable With Rising Rentals?' (14 Jul 2026): "
      "rent 10–12% of revenue recommended (8–15% typical); food 28–35%; staff 18–25%; utilities 4–7%; "
      "marketing 2–5%; EBITDA 8–18%")
S4 = ("S4: Petpooja, 'Free Restaurant P&L Template' (updated May 2026): food cost 28–35%, labour 20–25%, "
      "rent under 12%, EBITDA 15–20%, delivery-app commission 22–25%. Page calls these "
      "'directional, not a published standard'")
SOURCES = [S1, S2, S3, S4]

# ---------------------------------------------------------------------------------------------------
# Inputs: key, label, value, unit, low, high, source/basis, assumption ID
# Money is in ₹ lakh (1 lakh = ₹100,000) throughout, so figures read the way founders quote them.
# ---------------------------------------------------------------------------------------------------
SECTIONS = [
    ("Scenario", [
        ("scenario", "Selected scenario: steady-state monthly gross sales per outlet", 8, "₹ lakh / month",
         6, 10, "Brief's ₹6L / ₹8L / ₹10L, read as monthly gross sales per outlet, ex-GST. Pick 6, 8 or 10.", "A3"),
    ]),
    ("Outlet set-up (one-time)", [
        ("area", "Outlet size", 600, "sq ft", 400, 700, "Café-QSR format; S1 lists 400–1,000 sq ft café formats", "A2"),
        ("capex", "Capex: fit-out, kitchen and coffee equipment, signage", 24, "₹ lakh", 15, 40,
         "Derived: S1/S2 total investment ₹20–50L less fee and working capital. Estimate", "A2"),
        ("fee", "Franchise fee", 6, "₹ lakh", 5, 10, "S1: ₹5–8L (Barista), ₹10L (BrewBakes)", "A2"),
        ("wc", "Working capital: opening stock, deposits, pre-opening costs", 2, "₹ lakh", 1, 4,
         "Estimate: about 3 weeks of fixed costs", "A2"),
    ]),
    ("Ramp-up (% of steady-state sales)", [
        ("ramp1", "Month 1", 0.60, "%", 0.40, 0.80, "Estimate: new-outlet ramp; month 7 onwards = 100%", "A2"),
        ("ramp2", "Month 2", 0.70, "%", 0.50, 0.85, "Estimate", "A2"),
        ("ramp3", "Month 3", 0.80, "%", 0.60, 0.90, "Estimate", "A2"),
        ("ramp4", "Month 4", 0.90, "%", 0.70, 0.95, "Estimate", "A2"),
        ("ramp5", "Month 5", 0.95, "%", 0.80, 1.00, "Estimate", "A2"),
        ("ramp6", "Month 6", 1.00, "%", 0.90, 1.00, "Estimate", "A2"),
    ]),
    ("Variable costs (% of gross sales)", [
        ("cogs", "COGS: food, beverage and packaging", 0.30, "% of sales", 0.28, 0.35, "S3, S4: 28–35%", "A2"),
        ("deliv_share", "Share of sales through delivery apps", 0.30, "% of sales", 0.20, 0.40,
         "Estimate. No sourced figure for café-QSR found", "A2"),
        ("deliv_comm", "Delivery-app commission on those sales", 0.22, "% of delivery sales", 0.22, 0.25,
         "S4: 22–25%", "A2"),
        ("royalty", "Royalty to the brand", 0.06, "% of sales", 0.04, 0.08,
         "S1: 5–8%; S2: ~4–6%", "A2"),
        ("mktg", "Marketing / brand-fund fee", 0.02, "% of sales", 0.02, 0.05, "S3: marketing 2–5%", "A2"),
        ("other_var", "Other variable opex: POS, card fees, consumables, repairs", 0.03, "% of sales", 0.02, 0.05,
         "Estimate", "A2"),
    ]),
    ("Fixed monthly costs", [
        ("rent", "Rent incl. maintenance", 0.90, "₹ lakh / month", 0.64, 1.20,
         "Set at 11% of ₹8L base sales; S3 recommends 10–12% (8–15% typical)", "A2"),
        ("staff", "Staff: salaries, food, incentives", 1.60, "₹ lakh / month", 1.44, 2.00,
         "Set at 20% of ₹8L base sales; S3 18–25%, S4 20–25%", "A2"),
        ("utilities", "Utilities: power, water, gas, internet", 0.40, "₹ lakh / month", 0.32, 0.56,
         "Set at 5% of ₹8L base sales; S3 4–7%", "A2"),
    ]),
    ("Scale100x commercials (per brand). Scale100x pricing is not public: all editable estimates", [
        ("setup", "One-time brand setup fee (Year 1)", 5, "₹ lakh", 2, 10, "Estimate. Pricing not public", "A1"),
        ("retainer", "Monthly retainer", 1.5, "₹ lakh / month", 1, 3, "Estimate. Pricing not public", "A1"),
        ("months", "Retainer months billed per year", 12, "months", 6, 12, "Estimate", "A1"),
        ("success_pct", "Success fee per franchisee signed, as % of the franchise fee", 0.25, "% of fee", 0.15, 0.50,
         "Estimate. Pricing not public", "A1"),
        ("roy_share", "Share of the brand's royalty paid to Scale100x (optional stream)", 0, "% of royalty", 0, 0.20,
         "Off by default. Set above 0 to model a royalty-share deal", "A1"),
    ]),
    ("FreshBrew roll-out and portfolio", [
        ("signed1", "FreshBrew outlets signed, Year 1", 24, "outlets", 12, 36,
         "Estimate: ramps to 3 a month, matching A19 volumes", "A27"),
        ("signed2", "FreshBrew outlets signed, Year 2", 36, "outlets", 24, 48, "Estimate: 3 a month (A19)", "A27"),
        ("signed3", "FreshBrew outlets signed, Year 3", 36, "outlets", 24, 48, "Estimate: 3 a month (A19)", "A27"),
        ("brands", "Brands in the portfolio scenario", 5, "brands", 1, 20,
         "Estimate. Scales one FreshBrew-like brand; Scale100x reports 20+ brands scaled", "A4"),
    ]),
]

INPUTS = {}
for _, rows in SECTIONS:
    for key, label, value, unit, low, high, source, aid in rows:
        INPUTS[key] = value

# ---------------------------------------------------------------------------------------------------
# Styles (inputs blue, calculations black, per the founder-friendly convention)
# ---------------------------------------------------------------------------------------------------
BLUE = Font(color="1F4FD8", bold=True)
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
H2 = Font(bold=True, size=11, color="FFFFFF")
NOTE = Font(italic=True, color="555555")
HEAD_FILL = PatternFill("solid", fgColor="111118")
SECTION_FILL = PatternFill("solid", fgColor="E8EEF9")
INPUT_FILL = PatternFill("solid", fgColor="FFF8E1")
THIN = Side(style="thin", color="CCCCCC")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
LAKH = '0.00'
PCT = '0.0%'
INT = '0'


def header(ws, row, labels, widths=None):
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=text)
        c.font = H2
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def build():
    wb = Workbook()

    # ------------------------------------------------------------------ Assumptions
    a = wb.active
    a.title = "Assumptions"
    a["A1"] = "Assumptions: every input lives here"
    a["A1"].font = TITLE
    a["A2"] = ("Blue cells on a cream background are inputs: change them and every tab updates. Black cells are "
               "formulas. Money is in ₹ lakh (₹1 lakh = ₹100,000). Revenue excludes GST; payback is pre-tax (A28).")
    a["A2"].font = NOTE
    a.merge_cells("A2:H2")
    a.row_dimensions[2].height = 30
    a["A2"].alignment = WRAP
    header(a, 4, ["Key", "Input", "Value", "Unit", "Low", "High", "Source / basis", "Assumption"],
           [12, 52, 11, 18, 9, 9, 70, 11])
    ref = {}  # key -> absolute reference like Assumptions!$C$6
    row = 5
    for title, rows in SECTIONS:
        c = a.cell(row=row, column=1, value=title)
        c.font = BOLD
        for col in range(1, 9):
            a.cell(row=row, column=col).fill = SECTION_FILL
        row += 1
        for key, label, value, unit, low, high, source, aid in rows:
            fmt = PCT if unit.startswith("%") else (INT if unit in ("sq ft", "outlets", "brands", "months") else LAKH)
            a.cell(row=row, column=1, value=key)
            a.cell(row=row, column=2, value=label)
            v = a.cell(row=row, column=3, value=value)
            v.font = BLUE
            v.fill = INPUT_FILL
            v.number_format = fmt
            a.cell(row=row, column=4, value=unit)
            for col, val in ((5, low), (6, high)):
                cell = a.cell(row=row, column=col, value=val)
                cell.number_format = fmt
            s = a.cell(row=row, column=7, value=source)
            s.alignment = WRAP
            a.cell(row=row, column=8, value=aid)
            ref[key] = f"Assumptions!$C${row}"
            row += 1
    dv = DataValidation(type="list", formula1='"6,8,10"', allow_blank=False)
    a.add_data_validation(dv)
    dv.add(ref["scenario"].split("!")[1].replace("$", ""))

    # Calculated helpers (black): the cost structure the Payback and Sensitivity tabs reuse.
    row += 1
    a.cell(row=row, column=1, value="Calculated from the inputs above (do not edit)").font = BOLD
    for col in range(1, 9):
        a.cell(row=row, column=col).fill = SECTION_FILL
    row += 1
    calc = [
        ("invest", "Initial investment = capex + franchise fee + working capital",
         f"={ref['capex']}+{ref['fee']}+{ref['wc']}", "₹ lakh", LAKH),
        ("var_pct", "Total variable costs as % of sales",
         f"={ref['cogs']}+{ref['deliv_share']}*{ref['deliv_comm']}+{ref['royalty']}+{ref['mktg']}+{ref['other_var']}",
         "% of sales", PCT),
        ("fixed", "Total fixed costs per month", f"={ref['rent']}+{ref['staff']}+{ref['utilities']}",
         "₹ lakh / month", LAKH),
        ("breakeven", "Break-even monthly sales = fixed costs ÷ (1 − variable %)",
         None, "₹ lakh / month", LAKH),
        ("success_fee", "Success fee per franchisee signed", f"={ref['success_pct']}*{ref['fee']}", "₹ lakh", LAKH),
    ]
    for key, label, formula, unit, fmt in calc:
        if key == "breakeven":
            formula = f"={ref['fixed']}/(1-{ref['var_pct']})"
        a.cell(row=row, column=1, value=key)
        a.cell(row=row, column=2, value=label)
        cell = a.cell(row=row, column=3, value=formula)
        cell.number_format = fmt
        a.cell(row=row, column=4, value=unit)
        ref[key] = f"Assumptions!$C${row}"
        row += 1

    row += 1
    a.cell(row=row, column=1, value="Sources (third-party listings and trade articles, not franchisor disclosures; "
                                    f"{ACCESSED})").font = BOLD
    row += 1
    for s in SOURCES:
        a.cell(row=row, column=1, value=s).alignment = WRAP
        a.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        a.row_dimensions[row].height = 32
        row += 1
    a.freeze_panes = "A5"

    # ------------------------------------------------------------------ P&L_Y1
    p = wb.create_sheet("P&L_Y1")
    p["A1"] = "Franchisee P&L, Year 1, selected scenario (₹ lakh)"
    p["A1"].font = TITLE
    p["A2"] = "=\"Scenario: ₹\"&" + ref["scenario"] + "&\"L steady-state monthly sales per outlet. Change it on Assumptions.\""
    p["A2"].font = NOTE
    header(p, 4, ["Line item"] + [f"M{m}" for m in range(1, 13)] + ["Year 1"], [44] + [9] * 12 + [11])
    lines = [
        ("month", "Month number"),
        ("ramp", "Ramp-up (% of steady state)"),
        ("rev", "Revenue: gross sales, ex-GST"),
        ("cogs", "COGS"),
        ("gp", "Gross profit"),
        ("deliv", "Delivery-app commission"),
        ("royalty", "Royalty"),
        ("mktg", "Marketing / brand fund"),
        ("other", "Other variable opex"),
        ("rent", "Rent"),
        ("staff", "Staff"),
        ("utilities", "Utilities"),
        ("opex", "Total operating costs (below gross profit)"),
        ("ebitda", "EBITDA"),
        ("margin", "EBITDA margin"),
        ("cum", "Cumulative cash (after initial investment)"),
    ]
    prow = {k: 5 + i for i, (k, _) in enumerate(lines)}
    ramp_refs = [ref[f"ramp{i}"] for i in range(1, 7)]
    for k, label in lines:
        r = prow[k]
        p.cell(row=r, column=1, value=label).font = BOLD if k in ("rev", "gp", "ebitda", "cum") else Font()
        for m in range(1, 13):
            col = get_column_letter(m + 1)
            prev = get_column_letter(m)

            def c(key):
                return f"{col}{prow[key]}"

            f = {
                "month": f"={m}",
                "ramp": (f"={ramp_refs[m - 1]}" if m <= 6 else "=1"),
                "rev": f"={ref['scenario']}*{c('ramp')}",
                "cogs": f"={c('rev')}*{ref['cogs']}",
                "gp": f"={c('rev')}-{c('cogs')}",
                "deliv": f"={c('rev')}*{ref['deliv_share']}*{ref['deliv_comm']}",
                "royalty": f"={c('rev')}*{ref['royalty']}",
                "mktg": f"={c('rev')}*{ref['mktg']}",
                "other": f"={c('rev')}*{ref['other_var']}",
                "rent": f"={ref['rent']}",
                "staff": f"={ref['staff']}",
                "utilities": f"={ref['utilities']}",
                "opex": f"=SUM({col}{prow['deliv']}:{col}{prow['utilities']})",
                "ebitda": f"={c('gp')}-{c('opex')}",
                "margin": f"=IF({c('rev')}=0,0,{c('ebitda')}/{c('rev')})",
                "cum": (f"=-{ref['invest']}+{c('ebitda')}" if m == 1 else f"={prev}{prow['cum']}+{c('ebitda')}"),
            }[k]
            cell = p.cell(row=r, column=m + 1, value=f)
            cell.number_format = PCT if k in ("ramp", "margin") else (INT if k == "month" else LAKH)
        total_col = "N"
        if k in ("rev", "cogs", "gp", "deliv", "royalty", "mktg", "other", "rent", "staff", "utilities", "opex", "ebitda"):
            p[f"{total_col}{r}"] = f"=SUM(B{r}:M{r})"
            p[f"{total_col}{r}"].number_format = LAKH
            p[f"{total_col}{r}"].font = BOLD
        elif k == "margin":
            p[f"{total_col}{r}"] = f"=IF(N{prow['rev']}=0,0,N{prow['ebitda']}/N{prow['rev']})"
            p[f"{total_col}{r}"].number_format = PCT
        elif k == "cum":
            p[f"{total_col}{r}"] = f"=M{r}"
            p[f"{total_col}{r}"].number_format = LAKH
    p.freeze_panes = "B5"

    # ------------------------------------------------------------------ Payback (36 months, all scenarios)
    b = wb.create_sheet("Payback")
    b["A1"] = "Payback: 36-month cash flow for the selected scenario and for ₹6L / ₹8L / ₹10L (₹ lakh)"
    b["A1"].font = TITLE
    b["A2"] = ("Monthly EBITDA = sales × (1 − variable %) − fixed costs. Cumulative cash starts at minus the initial "
               "investment. Payback month = the first month cumulative cash is 0 or more (MATCH on the 'Paid back?' "
               "column). Pre-tax, ex-GST. Sales are flat after ramp-up (no growth assumed).")
    b["A2"].font = NOTE
    b["A2"].alignment = WRAP
    b.merge_cells("A2:N2")
    b.row_dimensions[2].height = 45
    groups = [("Selected", ref["scenario"]), ("₹6L", "6"), ("₹8L", "8"), ("₹10L", "10")]
    # Scenario sales levels sit in row 4 as plain labels-with-values; 6/8/10 are the brief's own scenarios.
    b["A4"] = "Steady-state monthly sales (₹ lakh)"
    b["A4"].font = BOLD
    hdr = ["Month", "Ramp"]
    for gi, (name, src) in enumerate(groups):
        col = 3 + gi * 3
        cell = b.cell(row=4, column=col, value=(f"={src}" if gi == 0 else int(src)))
        cell.number_format = LAKH
        cell.font = BOLD
        hdr += [f"{name} sales", f"{name} EBITDA", f"{name} cumulative"]
    hdr += ["Selected paid back?", "₹6L paid back?", "₹8L paid back?", "₹10L paid back?"]
    header(b, 5, hdr, [8, 8] + [11] * 12 + [12] * 4)
    first, last = 6, 41
    for m in range(1, 37):
        r = 5 + m
        b.cell(row=r, column=1, value=m)
        b.cell(row=r, column=2, value=(f"={ramp_refs[m - 1]}" if m <= 6 else "=1")).number_format = PCT
        for gi in range(4):
            sc = 3 + gi * 3
            s_col, e_col, c_col = (get_column_letter(sc + i) for i in range(3))
            level = f"{s_col}$4"
            b.cell(row=r, column=sc, value=f"={level}*$B{r}").number_format = LAKH
            b.cell(row=r, column=sc + 1,
                   value=f"={s_col}{r}*(1-{ref['var_pct']})-{ref['fixed']}").number_format = LAKH
            prev = f"-{ref['invest']}" if m == 1 else f"{c_col}{r - 1}"
            b.cell(row=r, column=sc + 2, value=f"={prev}+{e_col}{r}").number_format = LAKH
            b.cell(row=r, column=15 + gi, value=f"={c_col}{r}>=0")
    paycols = {}
    b["A43"] = "Payback month (first month cumulative ≥ 0)"
    b["A43"].font = BOLD
    for gi, (name, _) in enumerate(groups):
        flag_col = get_column_letter(15 + gi)
        out_col = get_column_letter(3 + gi * 3)
        cell = b[f"{out_col}43"]
        cell.value = f'=IFERROR(MATCH(TRUE,{flag_col}{first}:{flag_col}{last},0),"> 36")'
        cell.font = BOLD
        paycols[name] = f"Payback!${out_col}$43"
    b["A44"] = "Cumulative cash at month 36"
    for gi in range(4):
        col = get_column_letter(5 + gi * 3)
        b[f"{get_column_letter(3 + gi * 3)}44"] = f"={col}{last}"
        b[f"{get_column_letter(3 + gi * 3)}44"].number_format = LAKH
    b.freeze_panes = "C6"

    # ------------------------------------------------------------------ Sensitivity
    s = wb.create_sheet("Sensitivity")
    s["A1"] = "Sensitivity: ₹6L / ₹8L / ₹10L steady-state monthly sales per outlet"
    s["A1"].font = TITLE
    s["A2"] = ("Same costs in every column; only sales change. Each figure is a direct formula on the Payback tab, "
               "not a pasted value.")
    s["A2"].font = NOTE
    header(s, 4, ["Metric", "₹6L", "₹8L (base)", "₹10L"], [46, 13, 13, 13])
    metrics = [
        ("Steady-state monthly sales (₹ lakh)", lambda c: f"=Payback!{c}$4", LAKH),
        ("Year 1 sales (₹ lakh)", lambda c: f"=SUM(Payback!{c}6:{c}17)", LAKH),
        ("Year 1 EBITDA (₹ lakh)", lambda c: f"=SUM(Payback!{chr(ord(c) + 1)}6:{chr(ord(c) + 1)}17)", LAKH),
        ("Year 1 EBITDA margin", None, PCT),
        ("Steady-state monthly EBITDA (₹ lakh)", lambda c: f"=Payback!{chr(ord(c) + 1)}41", LAKH),
        ("Steady-state EBITDA margin", None, PCT),
        ("Payback month", lambda c: f"=Payback!{c}43", INT),
        ("Cumulative cash at month 36 (₹ lakh)", lambda c: f"=Payback!{c}44", LAKH),
    ]
    scen_cols = {"B": "F", "C": "I", "D": "L"}  # Sensitivity column -> Payback sales column
    for i, (label, fn, fmt) in enumerate(metrics):
        r = 5 + i
        s.cell(row=r, column=1, value=label)
        for scol, pcol in scen_cols.items():
            if label == "Year 1 EBITDA margin":
                f = f"=IF({scol}6=0,0,{scol}7/{scol}6)"
            elif label == "Steady-state EBITDA margin":
                f = f"=IF({scol}5=0,0,{scol}9/{scol}5)"
            else:
                f = fn(pcol)
            cell = s[f"{scol}{r}"]
            cell.value = f
            cell.number_format = fmt
    s["A14"] = "Break-even monthly sales, same in all three (₹ lakh)"
    s["A14"].font = BOLD
    s["B14"] = f"={ref['breakeven']}"
    s["B14"].number_format = LAKH
    s["A15"] = "Initial investment (₹ lakh)"
    s["B15"] = f"={ref['invest']}"
    s["B15"].number_format = LAKH

    # ------------------------------------------------------------------ Scale100x_Rev
    x = wb.create_sheet("Scale100x_Rev")
    x["A1"] = "Scale100x revenue from the FreshBrew engagement, Years 1–3 (₹ lakh)"
    x["A1"].font = TITLE
    x["A2"] = ("Kept separate from franchisee economics. Scale100x's pricing is not public, so every commercial input "
               "is an editable estimate (A1).")
    x["A2"].font = NOTE
    header(x, 4, ["Line", "Year 1", "Year 2", "Year 3", "3-year total"], [52, 12, 12, 12, 14])
    xrows = [
        ("year", "Year number"),
        ("signed", "FreshBrew outlets signed"),
        ("cumulative", "Outlets signed to date (end of year)"),
        ("setup", "Setup fee (Year 1 only)"),
        ("retainer", "Retainer"),
        ("success", "Success fees"),
        ("royshare", "Royalty share (optional, off by default)"),
        ("total", "Scale100x revenue from FreshBrew"),
        ("portfolio", "Portfolio scenario: × brands like FreshBrew"),
    ]
    xr = {k: 5 + i for i, (k, _) in enumerate(xrows)}
    signed_refs = [ref["signed1"], ref["signed2"], ref["signed3"]]
    for k, label in xrows:
        r = xr[k]
        x.cell(row=r, column=1, value=label).font = BOLD if k in ("total", "portfolio") else Font()
        for y in range(1, 4):
            col = get_column_letter(y + 1)
            prev = get_column_letter(y)
            f = {
                "year": f"={y}",
                "signed": f"={signed_refs[y - 1]}",
                "cumulative": (f"={col}{xr['signed']}" if y == 1 else f"={prev}{xr['cumulative']}+{col}{xr['signed']}"),
                "setup": f"=IF({col}{xr['year']}=1,{ref['setup']},0)",
                "retainer": f"={ref['retainer']}*{ref['months']}",
                "success": f"={col}{xr['signed']}*{ref['success_fee']}",
                # Average outlets trading in the year × selected monthly sales × 12 × royalty % × Scale100x share.
                "royshare": (f"=(({col}{xr['cumulative']}-{col}{xr['signed']})+{col}{xr['cumulative']})/2"
                             f"*{ref['scenario']}*12*{ref['royalty']}*{ref['roy_share']}"),
                "total": f"=SUM({col}{xr['setup']}:{col}{xr['royshare']})",
                "portfolio": f"={col}{xr['total']}*{ref['brands']}",
            }[k]
            cell = x.cell(row=r, column=y + 1, value=f)
            cell.number_format = INT if k in ("year", "signed", "cumulative") else LAKH
        if k in ("signed", "setup", "retainer", "success", "royshare", "total", "portfolio"):
            x[f"E{r}"] = f"=SUM(B{r}:D{r})"
            x[f"E{r}"].number_format = INT if k == "signed" else LAKH
            x[f"E{r}"].font = BOLD
    x["A15"] = "Revenue mix, 3 years"
    x["A15"].font = BOLD
    for i, k in enumerate(("setup", "retainer", "success", "royshare")):
        r = 16 + i
        x[f"A{r}"] = f"=A{xr[k]}"
        x[f"B{r}"] = f"=IF($E${xr['total']}=0,0,E{xr[k]}/$E${xr['total']})"
        x[f"B{r}"].number_format = PCT
    x["A21"] = ("Note: the portfolio row multiplies one FreshBrew-like engagement by the number of brands. It is a "
                "scale illustration, not a forecast.")
    x["A21"].font = NOTE

    # ------------------------------------------------------------------ Checks
    k = wb.create_sheet("Checks")
    k["A1"] = "Checks: every row should read TRUE"
    k["A1"].font = TITLE
    header(k, 3, ["Check", "Result"], [96, 10])
    P = "'P&L_Y1'"
    checks = [
        ("P&L: every month, EBITDA = revenue − all cost lines",
         f"=SUMPRODUCT(ABS({P}!B{prow['ebitda']}:M{prow['ebitda']}-({P}!B{prow['rev']}:M{prow['rev']}"
         f"-{P}!B{prow['cogs']}:M{prow['cogs']}-{P}!B{prow['deliv']}:M{prow['deliv']}"
         f"-{P}!B{prow['royalty']}:M{prow['royalty']}-{P}!B{prow['mktg']}:M{prow['mktg']}"
         f"-{P}!B{prow['other']}:M{prow['other']}-{P}!B{prow['rent']}:M{prow['rent']}"
         f"-{P}!B{prow['staff']}:M{prow['staff']}-{P}!B{prow['utilities']}:M{prow['utilities']})))<0.0001"),
        ("P&L Year 1 EBITDA total = sum of the 12 months",
         f"=ABS({P}!N{prow['ebitda']}-SUM({P}!B{prow['ebitda']}:M{prow['ebitda']}))<0.0001"),
        ("P&L Year 1 EBITDA (built line by line) = Payback selected EBITDA months 1–12 (built from the "
         "variable % and fixed costs): two different routes, same answer",
         f"=ABS({P}!N{prow['ebitda']}-SUM(Payback!D6:D17))<0.0001"),
        ("Initial investment = capex + franchise fee + working capital",
         f"=ABS({ref['invest']}-({ref['capex']}+{ref['fee']}+{ref['wc']}))<0.0001"),
        ("P&L month-12 cumulative cash = Payback month-12 cumulative (selected)",
         f"=ABS({P}!M{prow['cum']}-Payback!E17)<0.0001"),
        ("Payback month (selected) is the first month cumulative cash ≥ 0, and the month before is still < 0",
         '=IF(ISNUMBER(Payback!C43),AND(INDEX(Payback!E6:E41,Payback!C43)>=0,'
         'IF(Payback!C43=1,TRUE,INDEX(Payback!E6:E41,Payback!C43-1)<0)),Payback!E41<0)'),
        ("Payback month for ₹8L in Sensitivity matches the selected scenario when ₹8L is selected",
         f'=IF({ref["scenario"]}=8,Sensitivity!C11=Payback!C43,TRUE)'),
        ("Scale100x: the four streams add up to the yearly total (Years 1–3)",
         f"=SUMPRODUCT(ABS(Scale100x_Rev!B{xr['total']}:D{xr['total']}-(Scale100x_Rev!B{xr['setup']}:D{xr['setup']}"
         f"+Scale100x_Rev!B{xr['retainer']}:D{xr['retainer']}+Scale100x_Rev!B{xr['success']}:D{xr['success']}"
         f"+Scale100x_Rev!B{xr['royshare']}:D{xr['royshare']})))<0.0001"),
        ("Scale100x: 3-year total = Year 1 + Year 2 + Year 3",
         f"=ABS(Scale100x_Rev!E{xr['total']}-SUM(Scale100x_Rev!B{xr['total']}:D{xr['total']}))<0.0001"),
        ("Selected scenario is one of the brief's three (6, 8 or 10)",
         f"=OR({ref['scenario']}=6,{ref['scenario']}=8,{ref['scenario']}=10)"),
    ]
    for i, (label, formula) in enumerate(checks):
        r = 4 + i
        k.cell(row=r, column=1, value=label)
        k.cell(row=r, column=2, value=formula)
    all_row = 4 + len(checks)
    k.cell(row=all_row, column=1, value="All checks pass").font = BOLD
    k.cell(row=all_row, column=2, value=f"=AND(B4:B{all_row - 1})").font = BOLD

    # Hand spot-checks: worked by hand at base inputs (₹8L selected), compared to the model.
    r = all_row + 2
    k.cell(row=r, column=1, value="Hand spot-checks at the base inputs (₹8L selected). Expected values were worked "
                                  "out by hand, outside the sheet.").font = BOLD
    spot = [
        ("1. Month-1 revenue = 8 × 60% = 4.80", f"={P}!B{prow['rev']}", 4.80),
        ("2. Steady-state monthly EBITDA = 8 × (1 − 0.476) − 2.90 = 1.292, where variable % = 30% + 30%×22% + "
         "6% + 2% + 3% = 47.6% and fixed = 0.90 + 1.60 + 0.40 = 2.90", "=Payback!D41", 1.292),
        ("3. Year 1 Scale100x revenue = 5 setup + 1.5 × 12 retainer + 24 × (25% × 6) success = 59.00",
         f"=Scale100x_Rev!B{xr['total']}", 59.0),
    ]
    k.cell(row=r + 1, column=1, value="Calculation")
    k.cell(row=r + 1, column=2, value="Model")
    k.cell(row=r + 1, column=3, value="By hand")
    k.cell(row=r + 1, column=4, value="Match?")
    k.column_dimensions["C"].width = 10
    k.column_dimensions["D"].width = 10
    for i, (label, formula, expected) in enumerate(spot):
        rr = r + 2 + i
        k.cell(row=rr, column=1, value=label).alignment = WRAP
        k.cell(row=rr, column=2, value=formula).number_format = '0.000'
        k.cell(row=rr, column=3, value=expected).number_format = '0.000'
        k.cell(row=rr, column=4, value=f'=IF({ref["scenario"]}=8,ABS(B{rr}-C{rr})<0.0005,"select ₹8L")')
    k.cell(row=r + 6, column=1, value="Manual check: no hard-coded numbers inside formulas except 0, 1, month and "
                                      "year indexes, 12 months a year, and the brief's 6 / 8 / 10 scenario levels. "
                                      "Ticked after a formula review on 17 Oct 2026.").font = NOTE

    # ------------------------------------------------------------------ Summary (first tab)
    m = wb.create_sheet("Summary", 0)
    m["A1"] = "FreshBrew QSR franchise model: summary"
    m["A1"].font = TITLE
    m["A2"] = ("How to use: change any blue input on Assumptions (including the ₹6L / ₹8L / ₹10L scenario) and "
               "every tab updates. ₹6L / ₹8L / ₹10L = steady-state monthly gross sales per outlet, ex-GST (A3). "
               "Money in ₹ lakh. Payback is pre-tax (A28). Inputs are sourced ranges or labelled estimates; "
               "Scale100x pricing is not public, so its fees are editable estimates (A1).")
    m["A2"].font = NOTE
    m["A2"].alignment = WRAP
    m.merge_cells("A2:F2")
    m.row_dimensions[2].height = 60
    header(m, 4, ["Franchisee economics (one outlet)", "Value", "", "Scale100x revenue (FreshBrew)", "Value", ""],
           [46, 13, 3, 40, 13, 3])
    left = [
        ("Selected scenario: monthly sales (₹ lakh)", f"={ref['scenario']}", LAKH),
        ("Initial investment (₹ lakh)", f"={ref['invest']}", LAKH),
        ("Payback month", "=Payback!C43", INT),
        ("Year 1 sales (₹ lakh)", f"={P}!N{prow['rev']}", LAKH),
        ("Year 1 EBITDA (₹ lakh)", f"={P}!N{prow['ebitda']}", LAKH),
        ("Steady-state monthly EBITDA margin", f"={P}!M{prow['margin']}", PCT),
        ("Break-even monthly sales (₹ lakh)", f"={ref['breakeven']}", LAKH),
    ]
    right = [
        ("Year 1 (₹ lakh)", f"=Scale100x_Rev!B{xr['total']}", LAKH),
        ("Year 2 (₹ lakh)", f"=Scale100x_Rev!C{xr['total']}", LAKH),
        ("Year 3 (₹ lakh)", f"=Scale100x_Rev!D{xr['total']}", LAKH),
        ("3-year total (₹ lakh)", f"=Scale100x_Rev!E{xr['total']}", LAKH),
        ("FreshBrew outlets signed, 3 years", f"=Scale100x_Rev!E{xr['signed']}", INT),
        ("Portfolio scenario, 3 years (₹ lakh)", f"=Scale100x_Rev!E{xr['portfolio']}", LAKH),
        ("All model checks pass", f"=Checks!B{all_row}", None),
    ]
    for i, (label, formula, fmt) in enumerate(left):
        m.cell(row=5 + i, column=1, value=label)
        c = m.cell(row=5 + i, column=2, value=formula)
        c.font = BOLD
        if fmt:
            c.number_format = fmt
    for i, (label, formula, fmt) in enumerate(right):
        m.cell(row=5 + i, column=4, value=label)
        c = m.cell(row=5 + i, column=5, value=formula)
        c.font = BOLD
        if fmt:
            c.number_format = fmt

    header(m, 13, ["Sensitivity", "₹6L", "₹8L", "₹10L"], None)
    mini = [("Steady-state monthly EBITDA (₹ lakh)", 9, LAKH), ("Steady-state EBITDA margin", 10, PCT),
            ("Year 1 EBITDA (₹ lakh)", 7, LAKH), ("Payback month", 11, INT)]
    for i, (label, srow, fmt) in enumerate(mini):
        r = 14 + i
        m.cell(row=r, column=1, value=label)
        for j, col in enumerate("BCD"):
            c = m.cell(row=r, column=2 + j, value=f"=Sensitivity!{col}{srow}")
            c.number_format = fmt

    # Charts
    cash = LineChart()
    cash.title = "Cumulative cash per outlet, 36 months (₹ lakh)"
    cash.y_axis.title = "₹ lakh"
    cash.x_axis.title = "Month"
    for col in (5, 8, 11, 14):
        cash.add_data(Reference(b, min_col=col, min_row=5, max_row=41), titles_from_data=True)
    cash.set_categories(Reference(b, min_col=1, min_row=6, max_row=41))
    cash.height, cash.width = 8, 18
    m.add_chart(cash, "A20")

    mix = BarChart()
    mix.type = "col"
    mix.grouping = "stacked"
    mix.overlap = 100
    mix.title = "Scale100x revenue by stream (₹ lakh)"
    mix.add_data(Reference(x, min_col=1, max_col=4, min_row=xr["setup"], max_row=xr["royshare"]),
                 from_rows=True, titles_from_data=True)
    mix.set_categories(Reference(x, min_col=2, max_col=4, min_row=4))
    mix.height, mix.width = 8, 14
    m.add_chart(mix, "D20")

    wb.save(OUT)
    return OUT


def expected():
    """Independent re-calculation in plain Python, used to check the sheet's formulas."""
    i = INPUTS
    var = i["cogs"] + i["deliv_share"] * i["deliv_comm"] + i["royalty"] + i["mktg"] + i["other_var"]
    fixed = i["rent"] + i["staff"] + i["utilities"]
    invest = i["capex"] + i["fee"] + i["wc"]
    ramp = [i[f"ramp{n}"] for n in range(1, 7)] + [1] * 30
    out = {"var_pct": var, "fixed": fixed, "invest": invest, "breakeven": fixed / (1 - var)}
    for level in (6, 8, 10):
        cum, payback, y1rev, y1e = -invest, None, 0, 0
        for mth, rp in enumerate(ramp, start=1):
            rev = level * rp
            e = rev * (1 - var) - fixed
            cum += e
            if mth <= 12:
                y1rev += rev
                y1e += e
            if payback is None and cum >= 0:
                payback = mth
        out[level] = {"y1_sales": y1rev, "y1_ebitda": y1e, "steady_ebitda": level * (1 - var) - fixed,
                      "payback": payback or "> 36", "cum36": cum}
    sf = i["success_pct"] * i["fee"]
    yrs = []
    for y, signed in enumerate((i["signed1"], i["signed2"], i["signed3"]), start=1):
        yrs.append((i["setup"] if y == 1 else 0) + i["retainer"] * i["months"] + signed * sf)
    out["scale100x"] = yrs
    out["scale100x_total"] = sum(yrs)
    out["portfolio"] = sum(yrs) * i["brands"]
    return out


if __name__ == "__main__":
    if "--expect" in sys.argv:
        import json
        print(json.dumps(expected(), indent=2, default=str))
    else:
        print(build())
